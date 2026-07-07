"""
Builds the 'Omni OS — Data Sources & Planning Phases.xlsx' reference workbook,
pulling the five-tenet, journey-stage, touchpoint and data-source-by-phase tables
straight out of the Campaign Planning & Strategy Deep Dive doc, plus a live tab
showing what's actually sitting in the local knowledge repository right now.
"""
from __future__ import annotations

import pathlib
import sqlite3

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = pathlib.Path(__file__).resolve().parent.parent
VAULT_DIR = BASE_DIR.parent
OUT_PATH = VAULT_DIR / "Omni OS — Data Sources & Planning Phases.xlsx"
DB_PATH = BASE_DIR / "data" / "omni_kb.db"

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
WRAP = Alignment(wrap_text=True, vertical="top")


def write_table(ws, headers: list[str], rows: list[list[str]], col_widths: list[int]):
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = WRAP
    for row in rows:
        ws.append(row)
    for r in range(2, len(rows) + 2):
        for c in range(1, len(headers) + 1):
            ws.cell(row=r, column=c).alignment = WRAP
    for idx, width in enumerate(col_widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A2"


def build_data_sources_sheet(wb):
    ws = wb.active
    ws.title = "Data Sources by Phase"
    headers = ["Planning Activity", "Public (free/open)", "Licensed & commercial (paid)", "Synthetic / derived"]
    rows = [
        ["Brand strategy & indication landscape",
         "FDA labels/DailyMed, ClinicalTrials.gov, PubMed/clinical literature, CDC/SEER epidemiology data, congress abstract archives (ASCO, ASH, AAN etc.), competitor pipeline disclosures (company IR filings)",
         "-",
         "-"],
        ["Segmentation & targeting",
         "CMS NPPES NPI registry (specialty, practice location), public HCO/hospital directories, publication authorship on PubMed (crude KOL proxy)",
         "OneKey (IQVIA) HCP master reference data; Rx/claims data - IQVIA, Komodo Health, Symphony Health (SHS); historical rep call activity from CRM (Veeva)",
         "Invisage-style digital affinity score, channel affinity score, market access index, Peer Impact Quotient (R3 model), propensity-to-prescribe/switch scores, microsegment/persona assignment (Stars/Potentials/Digital Enthusiasts/Conventionalists)"],
        ["Content plan & messaging",
         "PubMed literature, FDA label language, patient advocacy group content (unbranded disease education), Google Trends, social listening (Reddit, YouTube) for patient/HCP language",
         "Content categorization/tagging services (17-category content taxonomy in BioMarin's model), Seismic/PromoMats content libraries",
         "Content affinity score (which content categories an HCP engages with most), current-belief -> desired-belief message maps synthesized from rep notes + literature, GenAI auto-tagged claims/components libraries"],
        ["Channel mix & media planning",
         "Industry channel-consumption benchmarks (Pew Research, trade press), programmatic exchange contextual audience data",
         "Historical digital engagement logs (email opens, web sessions, webinar attendance) held in SFMC/CDP; media agency spend and performance data (FIA-governed data drops)",
         "MMx/marketing-mix-model response curves and ROI-by-channel outputs; next-best-channel/NBA recommendations; channel-preference and day/time-preference scores"],
        ["Omnichannel journey design",
         "Published HCP-journey/CX research (e.g., DT Consulting's State of Customer Experience reports)",
         "Indegene's own cross-program benchmark library (\"45+ omnichannel programs\")",
         "Journey-stage/adoption-ladder scoring models, simulated journey paths, illustrative synthetic personas"],
        ["Data & tech feasibility, launch pack creation",
         "Platform API/technical documentation (Veeva, SFMC, Salesforce Data Cloud) - feasibility input, not data",
         "Consent/preference data from CMP platforms (OneTrust); MDM golden-record HCP/account data; C360/P360 data repositories (Databricks)",
         "Auto-tagged/validated target lists, UTM taxonomy, DCR (data-change-request) and match/merge survivorship outputs"],
        ["Ongoing measurement & optimization",
         "-",
         "Channel performance data fed via FTP/S3 from media/channel partners under FIA agreements; claims/Rx refresh (Komodo/SHS/IQVIA)",
         "Engagement scores, NBA/badge assignments, MMx-driven reallocation recommendations, RCA-derived anomaly flags"],
    ]
    write_table(ws, headers, rows, [30, 45, 45, 45])


def build_five_tenets_sheet(wb):
    ws = wb.create_sheet("Five Tenets")
    headers = ["#", "Tenet", "The question it answers", "Where it shows up downstream"]
    rows = [
        [1, "Segmentation & persona", "Who are we talking to?", "Target lists, call plans, consent management"],
        [2, "Journey stage", "Where is this HCP on the path from unaware to champion?", "Which content/CTA is eligible to fire"],
        [3, "Messaging architecture", "What do we need them to believe next?", "Claims, content briefs, e-detail flow"],
        [4, "Channel mix / touchpoint plan", "How do we reach them, in what mix?", "Media plan, rep call plan, MMx spend split"],
        [5, "Measurement & feedback loop", "Did it move them up the journey?", "Engagement scoring, NBA, next planning cycle"],
    ]
    write_table(ws, headers, rows, [5, 30, 45, 40])


def build_journey_sheet(wb):
    ws = wb.create_sheet("HCP Journey Stages")
    headers = ["Stage", "HCP mental state", "Core barrier", "Engagement goal", "Typical messaging", "Primary touchpoints", "Signal that promotes them"]
    rows = [
        ["1. Unaware", "Never encountered the brand/molecule", "No awareness", "Get on the radar",
         "Unbranded disease-state / epidemiology",
         "Congress presence, programmatic/paid social, PubMed-adjacent placements, rep cold-call",
         "Any first impression/click/booth visit"],
        ["2. Aware", "Has heard of it, no real understanding", "\"I don't understand it\"", "Build comprehension",
         "Mechanism of action, how it differs from standard of care",
         "Email introducing the brand, banner/display, rep intro detail, unbranded-to-branded website",
         "Email open, site visit, first rep meeting accepted"],
        ["3. Interested / Evaluating", "Understands it, isn't yet convinced", "\"I don't believe it (for my patients)\"", "Build belief/credibility",
         "Efficacy & safety data, head-to-head evidence, KOL commentary",
         "Webinars, e-detailing, congress symposia, MSL scientific exchange, self-detail on web",
         "Webinar registration/attendance, 3+ min site visit, MSL request"],
        ["4. Trial / First Rx", "Convinced, hasn't acted yet", "How do I actually start", "Convert belief -> first prescription",
         "Patient selection criteria, dosing & initiation, access/reimbursement support",
         "Rep detail with starter/samples, patient support program enrollment, dosing app/CLM leave-behind",
         "First sample request, first patient-support enrollment, first Rx flagged in claims/Rx data"],
        ["5. Adoption / Regular prescriber", "Prescribes routinely", "Complacency or competitive switch risk", "Reinforce, defend share of mind",
         "Real-world evidence, broader patient-type expansion, practical troubleshooting (AE management)",
         "Ongoing rep cadence, nurture email flows, peer case-study content, targeted congress follow-up",
         "Repeat Rx / TRx trend, consistent digital engagement"],
        ["6. Advocate / Champion", "Prescribes routinely and influences peers", "Under-utilising them", "Turn into a voice for the brand",
         "Co-created content, real-world data they helped generate, peer-to-peer talking points",
         "Speaker programs, advisory boards, peer-to-peer/DOL programs, congress podium slots, co-authored abstracts",
         "Speaking engagements accepted, peer-referral pattern, advisory board participation"],
    ]
    write_table(ws, headers, rows, [22, 28, 24, 24, 34, 34, 30])


def build_touchpoints_sheet(wb):
    ws = wb.create_sheet("Touchpoints")
    headers = ["Category", "Touchpoints", "Best for", "Typical stage"]
    rows = [
        ["Reach", "Programmatic display, paid social, paid search, unbranded press/PR", "Getting on the radar at low cost per impression", "Unaware -> Aware"],
        ["Owned digital", "Branded email/nurture flows, CLM/e-detailing, brand website, portal/app", "Structured, sequenced comprehension-building", "Aware -> Interested"],
        ["Events", "Congress booths/symposia, webinars (ON24), speaker programs", "Concentrated, high-credibility bursts", "Interested -> Trial"],
        ["Field", "Rep details, samples, MSL scientific exchange", "Personalised, two-way, can close specific objections", "Interested -> Trial -> Adoption"],
        ["Peer", "KOL/DOL content, peer-to-peer programs, advisory boards, congress podium slots", "Credibility transfer between HCPs", "Adoption -> Champion"],
        ["Patient-adjacent", "Patient support program enrollment, EHR point-of-care alerts", "Practical initiation support, share-of-mind at the point of the Rx decision", "Trial -> Adoption"],
    ]
    write_table(ws, headers, rows, [18, 45, 45, 22])


def build_kb_stats_sheet(wb):
    ws = wb.create_sheet("Live Knowledge Repo Stats")
    headers = ["Source", "Search term", "Documents fetched"]
    rows = []
    if DB_PATH.exists():
        conn = sqlite3.connect(DB_PATH)
        cur = conn.execute(
            "SELECT source, search_term, COUNT(*) FROM documents GROUP BY source, search_term ORDER BY source, search_term"
        )
        rows = [list(r) for r in cur.fetchall()]
        conn.close()
    if not rows:
        rows = [["(no data yet - run scrapers/run_all.py first)", "-", 0]]
    write_table(ws, headers, rows, [22, 30, 18])
    ws.append([])
    total_row = ["TOTAL", "", sum(r[2] for r in rows if isinstance(r[2], int))]
    ws.append(total_row)


def main():
    wb = Workbook()
    build_data_sources_sheet(wb)
    build_five_tenets_sheet(wb)
    build_journey_sheet(wb)
    build_touchpoints_sheet(wb)
    build_kb_stats_sheet(wb)
    wb.save(OUT_PATH)
    print(f"Saved workbook to {OUT_PATH}")


if __name__ == "__main__":
    main()
