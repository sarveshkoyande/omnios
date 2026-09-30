"""The Campaign Planner's actions, called by the HTTP routes in app/server.py.

Every action that thinks (analyse, answer, review assumptions, generate or change the brief,
build the journey blueprint, deploy) runs as a background job (jobs.py) and streams events.
A job works on a copy of the session state and commits it only when it finishes; a failure
or a cancelled job leaves the session exactly as it was, so the step can simply be retried.

The state carries a conversation log (`items`) the Cockpit renders on the left, the Campaign
Briefing Document and the blueprint it renders on the right, and the inputs each later step
needs (the full requirements text -- including an uploaded document's -- the clarification
answers and the assumption decisions).
"""
from __future__ import annotations

import copy
import io
import threading
import time
import traceback
import uuid

from . import agents, flow_xml, jobs, llm, mermaid, salesforce, store
from . import briefing as bf

MAX_FILE_BYTES = 10 * 1024 * 1024  # Camille's upload limit
MAX_DOC_CHARS = 150_000
_EXTENSIONS = ("pdf", "docx", "txt", "md")
_REASONING_CAP = 80_000  # characters of each agent's streamed reasoning kept with the session

_LOCKS: dict[str, threading.Lock] = {}
_LOCKS_GUARD = threading.Lock()


def _lock(session_id: str) -> threading.Lock:
    with _LOCKS_GUARD:
        return _LOCKS.setdefault(session_id, threading.Lock())


def _now() -> str:
    return time.strftime("%Y-%m-%dT%H:%M:%SZ", time.gmtime())


def _iid() -> str:
    return uuid.uuid4().hex[:10]


def new_state(auto_assume: bool = False) -> dict:
    return {
        "schema": 1,
        "stage": "intake",            # intake | clarifying | assumptions | briefing | blueprint
        "auto_assume": auto_assume,
        "requirements": "",           # typed text + extracted document text
        "typed": "",
        "file_name": None,
        "is_document": False,
        "extracted_context": "",
        "context": "",                # requirements + clarifications (+ assumption decisions)
        "answers_text": "",
        "assumption_review": "",
        "questions": None,
        "assumptions": None,
        "briefing": None,
        "briefing_version": 0,
        "briefing_source": None,      # "ai" | "rules"
        "briefing_note": None,
        "blueprint": None,
        "blueprint_briefing_version": None,
        "blueprint_stale": False,
        "spec": None,
        "deployment": None,
        "items": [],
        "token_usage": 0,
    }


# ------------------------------------------------------------------ workspace -----------

def workspace(brand: str | None, plan_id: int | None, campaign_id: int | None) -> dict:
    """Names and brand-kit facts for the scope this session plans in."""
    ws = {"brand": brand, "plan_id": plan_id, "campaign_id": campaign_id, "plan_name": None, "campaign_name": None, "kit": {}}
    from strategy import brand_kit, hierarchy
    if plan_id:
        try:
            ws["plan_name"] = hierarchy.get_plan(plan_id).get("name")
        except Exception:  # noqa: BLE001 -- context only; a missing plan just isn't named
            pass
    if campaign_id:
        try:
            ws["campaign_name"] = hierarchy.get_campaign(campaign_id).get("name")
        except Exception:  # noqa: BLE001
            pass
    if brand:
        try:
            ws["kit"] = brand_kit.kit_for(brand) or {}
        except Exception:  # noqa: BLE001
            ws["kit"] = {}
    return ws


UNTITLED = "Untitled — Campaign plan"


def _default_title(ws: dict) -> str:
    return f"{ws['campaign_name']} — Campaign plan" if ws.get("campaign_name") else UNTITLED


def _name_untitled(session_id: str, b: dict) -> None:
    """An unscoped session stays "Untitled" until its first briefing names the campaign."""
    sess = store.get(session_id)
    name = str((b or {}).get("campaignName") or "").strip()
    if sess and sess["title"] == UNTITLED and name and name != bf.NOT_SPECIFIED:
        store.set_fields(session_id, title=f"{name} — Campaign plan")


def view(sess: dict) -> dict:
    job = jobs.active(sess["id"])
    return {
        "id": sess["id"], "brand": sess["brand"], "plan_id": sess["plan_id"], "campaign_id": sess["campaign_id"],
        "title": sess["title"], "artifact_id": sess["artifact_id"],
        "created_at": sess["created_at"], "updated_at": sess["updated_at"],
        "state": sess["state"],
        "job": {"id": job.id, "kind": job.kind} if job else None,
        "llm": {"available": llm.available(), "engine": llm.engine_label()},
    }


def _load(session_id: str) -> dict:
    sess = store.get(session_id)
    if not sess:
        raise KeyError(f"no campaign planner session {session_id}")
    return sess


def open_session(brand: str | None, plan_id: int | None, campaign_id: int | None) -> dict:
    sess = store.find(brand, campaign_id)
    if not sess:
        ws = workspace(brand, plan_id, campaign_id)
        sess = store.create(brand, plan_id, campaign_id, _default_title(ws), new_state())
    elif plan_id and sess["plan_id"] != plan_id:
        store.set_fields(sess["id"], plan_id=plan_id)  # the campaign moved to another plan
        sess = _load(sess["id"])
    return view(sess)


def session_for_artifact(artifact_id: str) -> dict | None:
    sess = store.find_by_artifact(artifact_id)
    return view(sess) if sess else None


def get_view(session_id: str) -> dict:
    return view(_load(session_id))


def rename(session_id: str, title: str) -> dict:
    _load(session_id)
    store.set_fields(session_id, title=title.strip() or UNTITLED)
    return get_view(session_id)


def _ensure_idle(session_id: str) -> None:
    if jobs.active(session_id):
        raise jobs.Busy("This campaign planner is already working on a step.")


# ------------------------------------------------------------------ documents -----------

def validate_upload(file_name: str, size: int) -> None:
    ext = file_name.rsplit(".", 1)[-1].lower() if "." in file_name else ""
    if ext not in _EXTENSIONS:
        raise ValueError("Unsupported file type. Please upload PDF, DOCX, TXT or MD files.")
    if size > MAX_FILE_BYTES:
        raise ValueError("File size exceeds 10MB limit")
    if size == 0:
        raise ValueError("The uploaded file is empty.")


def extract_document(file_name: str, content: bytes) -> str:
    """The document's text. DOCX includes tables in reading order (as Camille's mammoth
    extraction did); PDF and text go through the app's shared extractor."""
    ext = file_name.rsplit(".", 1)[-1].lower()
    if ext == "docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        doc = Document(io.BytesIO(content))
        lines: list[str] = []
        for child in doc.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]
            if tag == "p":
                lines.append(Paragraph(child, doc).text)
            elif tag == "tbl":
                for row in Table(child, doc).rows:
                    cells = [c.text.strip() for c in row.cells]
                    lines.append(" | ".join(c for c in cells if c))
        text = "\n".join(lines).strip()
        return text[:MAX_DOC_CHARS] + ("\n\n[...truncated...]" if len(text) > MAX_DOC_CHARS else "")
    from strategy.document_intake import extract_text
    return extract_text(file_name, content, max_chars=MAX_DOC_CHARS)


# ------------------------------------------------------------------ job plumbing --------

class _Run:
    """One job's working copy of the session, its events, and its commit."""

    def __init__(self, job: jobs.Job, sess: dict):
        self.job = job
        self.sid = sess["id"]
        self.sess = sess
        self.state = copy.deepcopy(sess["state"])
        self.ws = workspace(sess["brand"], sess["plan_id"], sess["campaign_id"])
        self.ws_block = bf.workspace_block(self.ws)
        self.pending_version: tuple[dict, str] | None = None
        self.new_version: int | None = None

    # -- events
    def emit(self, ev: dict) -> None:
        self.job.emit(ev)

    def stop(self) -> bool:
        return self.job.cancelled

    def item(self, item: dict) -> dict:
        item = {"id": _iid(), "at": _now(), **item}
        self.state["items"].append(item)
        self.emit({"type": "item", "item": item})
        return item

    def update_item(self, item: dict) -> None:
        self.emit({"type": "item_update", "item": item})

    def progress(self, item: dict, step: str, title: str, details: str) -> None:
        entry = {"step": step, "title": title, "details": details}
        item.setdefault("steps", []).append(entry)
        self.emit({"type": "progress", "item_id": item["id"], "step": entry})

    def tokens(self, usage: dict | None) -> int:
        n = int((usage or {}).get("totalTokenCount") or 0)
        self.state["token_usage"] = int(self.state.get("token_usage") or 0) + n
        return n

    def with_ws(self, text: str) -> str:
        return f"{text}\n\n{self.ws_block}" if self.ws_block else text

    # -- commit
    def commit(self) -> bool:
        """Save the working state (and a pending briefing version). False if cancelled."""
        with _lock(self.sid):
            if self.job.cancelled:
                return False
            if self.pending_version:
                b, reason = self.pending_version
                self.new_version = store.add_briefing_version(self.sid, b, reason)
                self.state["briefing_version"] = self.new_version
                _name_untitled(self.sid, b)
            store.save_state(self.sid, self.state)
        if self.pending_version:
            _sync_artifact(self.sid)
        return True

    def done(self, status: str = "ok") -> None:
        sess = store.get(self.sid)
        self.emit({"type": "done", "status": status, "session": view(sess) if sess else None})

    def fail(self, message: str) -> None:
        """End the job without saving: the session stays as it was before the job began."""
        self.emit({"type": "error", "message": message})
        self.done("error")


def _start(session_id: str, kind: str, body) -> dict:
    sess = _load(session_id)
    _ensure_idle(session_id)

    def work(job: jobs.Job) -> None:
        run = _Run(job, sess)
        try:
            body(run)
        except llm.Cancelled:
            run.emit({"type": "done", "status": "cancelled", "session": None})
        except Exception as exc:  # noqa: BLE001 -- surfaced to the user; the session is untouched
            if not isinstance(exc, ValueError):  # ValueErrors are the expected, user-facing ones
                traceback.print_exc()
            run.fail(str(exc) or exc.__class__.__name__)

    job = jobs.start(session_id, kind, work)
    return {"job_id": job.id, "kind": kind}


# ------------------------------------------------------------------ briefing ------------

def _answers_block(state: dict) -> str | None:
    parts = [p for p in (state.get("answers_text"), state.get("assumption_review")) if p]
    return "\n\n".join(parts) or None


def _generate_briefing(run: _Run, prog: dict, context: str, answers: str | None) -> None:
    """Generate the briefing (or assemble it by rules), then commit it as a new version."""
    state = run.state
    source, note = "ai", None
    b = None
    if llm.available():
        try:
            b, usage = bf.generate(run.with_ws(context), answers, should_stop=run.stop)
            run.tokens(usage)
        except llm.Cancelled:
            raise
        except Exception as exc:  # noqa: BLE001 -- OmniOS rule: degrade to rules, never hard-fail
            print(f"[campaign-creator] briefing generation failed, using rules: {exc}")
            source, note = "rules", f"The AI call failed ({exc}), so this briefing was assembled from your text by rules."
    else:
        source, note = "rules", "No AI model is configured, so this briefing was assembled from your text by rules."
    if b is None:
        b = bf.rules_briefing(context, run.ws)
    run.progress(prog, "complete", "Campaign Consultant", "Campaign Briefing Document generated successfully!")
    prog["status"] = "done"
    run.update_item(prog)
    state.update(briefing=b, briefing_source=source, briefing_note=note, stage="briefing",
                 questions=None, assumptions=None)
    if state.get("blueprint"):
        state["blueprint_stale"] = True
    run.pending_version = (b, "Generated")
    run.item({"kind": "agent", "text": "Your Campaign Briefing Document is ready. Review it, download it, request changes, "
                                        "or proceed to generate the journey blueprint."
                                        + (f"\n\n{note}" if note else "")})
    if run.commit():
        run.emit({"type": "briefing", "briefing": b, "version": run.new_version, "source": source, "note": note})
        run.done()


def analyze(session_id: str, typed: str, file_name: str | None, content: bytes | None, auto_assume: bool) -> dict:
    """Camille /campaign-briefing/generate: read the brief, ask or generate."""
    sess = _load(session_id)
    if sess["state"].get("stage") != "intake":
        raise ValueError("This campaign already has a brief. Start over to plan a new one.")
    typed = (typed or "").strip()
    if file_name and content is not None:
        validate_upload(file_name, len(content))
    elif not typed:
        raise ValueError("Please provide requirements or upload a document.")

    def body(run: _Run) -> None:
        state = run.state
        run.item({"kind": "user", "text": typed or f"[File uploaded: {file_name}]", "file_name": file_name})
        prog = run.item({"kind": "progress", "title": "Campaign Consultant", "status": "running", "steps": []})
        requirements, is_doc = typed, False
        if file_name and content is not None:
            run.progress(prog, "extracting", "Document Reader", f"Extracting content from {file_name}...")
            try:
                text = extract_document(file_name, content)
            except Exception as exc:  # noqa: BLE001
                raise ValueError(f"Failed to extract text from file: {exc}") from exc
            if not text.strip():
                raise ValueError("No text could be extracted from that file.")
            is_doc = True
            requirements = f"User Instructions: {typed}\n\nDocument Content:\n{text}" if typed else text
            run.progress(prog, "extracted", "Document Reader", f"Extracted {len(text)} characters successfully.")
        state.update(requirements=requirements, typed=typed, file_name=file_name, is_document=is_doc,
                     auto_assume=bool(auto_assume), context=requirements)
        run.progress(prog, "analyzing", "Campaign Consultant", "Analyzing requirements and identifying gaps...")
        questions: list[dict] = []
        if llm.available():
            try:
                analysis, usage = bf.analyze(requirements, is_doc, run.ws_block, should_stop=run.stop)
                run.tokens(usage)
            except llm.Cancelled:
                raise
            except Exception as exc:  # noqa: BLE001 -- as Camille: an unreadable analysis means "no questions"
                print(f"[campaign-creator] analysis failed, continuing without questions: {exc}")
                analysis = {"needsQuestions": False, "questions": [], "extractedContext": requirements}
            state["extracted_context"] = bf._text(analysis.get("extractedContext")) or requirements
            if analysis.get("needsQuestions"):
                questions = bf.normalize_questions(analysis.get("questions"))
        if questions and not auto_assume:
            run.progress(prog, "questions", "Campaign Consultant", "I need a few clarifications to create a comprehensive briefing...")
            prog["status"] = "done"
            run.update_item(prog)
            run.item({"kind": "questions", "questions": questions, "answered": None})
            state.update(stage="clarifying", questions=questions)
            if run.commit():
                run.done()
            return
        run.progress(prog, "generating", "Campaign Consultant", "Generating Campaign Briefing Document...")
        _generate_briefing(run, prog, requirements, None)

    return _start(session_id, "analyze", body)


def answer_questions(session_id: str, answers_text: str) -> dict:
    """Camille /campaign-briefing/answer: review the answers for open assumptions, else generate."""
    sess = _load(session_id)
    st = sess["state"]
    if st.get("stage") != "clarifying" or not st.get("questions"):
        raise ValueError("There are no clarification questions waiting for an answer.")
    answers_text = (answers_text or "").strip()
    if not answers_text:
        raise ValueError("Answer the questions first.")

    def body(run: _Run) -> None:
        state = run.state
        for it in reversed(state["items"]):
            if it.get("kind") == "questions" and it.get("answered") is None:
                it["answered"] = answers_text
                run.update_item(it)
                break
        run.item({"kind": "user", "text": answers_text})
        prog = run.item({"kind": "progress", "title": "Campaign Consultant", "status": "running", "steps": []})
        full = f"{state['requirements']}\n\nClarifications Provided by User:\n{answers_text}"
        state.update(answers_text=answers_text, context=full, questions=None)
        run.progress(prog, "analyzing_assumptions", "Campaign Consultant", "Reviewing for assumptions and ambiguities...")
        found: list[str] = []
        if llm.available():
            found, usage = bf.find_assumptions(run.with_ws(full), should_stop=run.stop)
            run.tokens(usage)
        if found:
            run.progress(prog, "assumptions_found", "Campaign Consultant", f"Found {len(found)} assumption(s) to review.")
            prog["status"] = "done"
            run.update_item(prog)
            run.item({"kind": "assumptions", "assumptions": found, "resolved": None})
            state.update(stage="assumptions", assumptions=found)
            if run.commit():
                run.done()
            return
        run.progress(prog, "generating", "Campaign Consultant",
                     "Incorporating your answers and generating Campaign Briefing Document...")
        _generate_briefing(run, prog, full, _answers_block(state))

    return _start(session_id, "answers", body)


def review_assumptions(session_id: str, mode: str, resolved_text: str = "") -> dict:
    """Camille's assumption checkpoint: accept them all, or override some, then generate."""
    sess = _load(session_id)
    st = sess["state"]
    if st.get("stage") != "assumptions" or not st.get("assumptions"):
        raise ValueError("There are no assumptions waiting for review.")
    if mode not in ("accept", "resolve"):
        raise ValueError("mode must be 'accept' or 'resolve'")
    resolved_text = (resolved_text or "").strip()
    if mode == "resolve" and not resolved_text:
        raise ValueError("Add a clarification, or accept the assumptions.")

    def body(run: _Run) -> None:
        state = run.state
        assumptions = state.get("assumptions") or []
        if mode == "accept":
            review = "\n\n".join(f"Ambiguity: {a}\nUser Clarification: Accepted default assumption" for a in assumptions)
            user_text, heading = "Accepted all detected assumptions.", "Assumptions Accepted by User"
        else:
            review = resolved_text
            user_text, heading = f"Resolved Ambiguities:\n{resolved_text}", "User Overrides on Assumptions"
        for it in reversed(state["items"]):
            if it.get("kind") == "assumptions" and it.get("resolved") is None:
                it["resolved"] = "Accepted all assumptions" if mode == "accept" else resolved_text
                run.update_item(it)
                break
        run.item({"kind": "user", "text": user_text})
        prog = run.item({"kind": "progress", "title": "Campaign Consultant", "status": "running", "steps": []})
        context = f"{state['context']}\n\n{heading}:\n{review}"
        state.update(assumption_review=review, context=context, assumptions=None)
        run.progress(prog, "generating", "Campaign Consultant", "Generating Campaign Briefing Document...")
        _generate_briefing(run, prog, context, _answers_block(state))

    return _start(session_id, "assumptions", body)


def update_briefing(session_id: str, modifications: str) -> dict:
    """Camille /campaign-briefing/update, on the briefing actually on screen."""
    sess = _load(session_id)
    if not sess["state"].get("briefing"):
        raise ValueError("There is no Campaign Briefing Document to change yet.")
    modifications = (modifications or "").strip()
    if not modifications:
        raise ValueError("Describe the change you want.")
    if not llm.available():
        raise ValueError("Changing the briefing needs the AI model, and none is configured on this server.")

    def body(run: _Run) -> None:
        state = run.state
        run.item({"kind": "user", "text": modifications})
        prog = run.item({"kind": "progress", "title": "Campaign Consultant", "status": "running", "steps": []})
        run.progress(prog, "updating", "Campaign Consultant", "Updating Campaign Briefing Document with your changes...")
        try:
            b, usage = bf.update(state["briefing"], modifications, should_stop=run.stop)
        except llm.Cancelled:
            raise
        except Exception as exc:  # noqa: BLE001
            print(f"[campaign-creator] briefing update failed: {exc}")
            raise ValueError("Failed to update Campaign Briefing. Please try again.") from exc
        run.tokens(usage)
        run.progress(prog, "complete", "Campaign Consultant", "Campaign Briefing Document updated successfully!")
        prog["status"] = "done"
        run.update_item(prog)
        state.update(briefing=b, briefing_source="ai", briefing_note=None)
        if state.get("blueprint"):
            state["blueprint_stale"] = True
        run.pending_version = (b, f"Changed: {modifications[:90]}")
        run.item({"kind": "agent", "text": "Updated the Campaign Briefing Document with your changes."})
        if run.commit():
            run.emit({"type": "briefing", "briefing": b, "version": run.new_version, "source": "ai", "note": None})
            run.done()

    return _start(session_id, "update", body)


def briefing_versions(session_id: str) -> list[dict]:
    _load(session_id)
    return store.briefing_versions(session_id)


def briefing_version(session_id: str, version: int) -> dict:
    _load(session_id)
    b = store.briefing_version(session_id, version)
    if b is None:
        raise KeyError(f"briefing version {version}")
    return {"version": version, "briefing": b}


def restore_briefing(session_id: str, version: int) -> dict:
    _ensure_idle(session_id)
    with _lock(session_id):
        sess = _load(session_id)
        b = store.briefing_version(session_id, version)
        if b is None:
            raise KeyError(f"briefing version {version}")
        state = sess["state"]
        new_v = store.add_briefing_version(session_id, b, f"Restored version {version}")
        state.update(briefing=b, briefing_version=new_v, briefing_source="ai" if state.get("briefing_source") == "ai" else state.get("briefing_source"))
        if state.get("blueprint"):
            state["blueprint_stale"] = True
        state["items"].append({"id": _iid(), "at": _now(), "kind": "agent", "text": f"Restored version {version} of the briefing (saved as v{new_v})."})
        store.save_state(session_id, state)
    _sync_artifact(session_id)
    return get_view(session_id)


# ------------------------------------------------------------------ blueprint -----------

def compile_prompt(state: dict, feedback: str | None = None) -> str:
    """What the blueprint agents read: the brief as given, the decisions taken on it, and the
    approved Campaign Briefing Document (Camille sent only the first of these)."""
    parts = ["Campaign Brief:", state.get("requirements") or ""]
    if state.get("answers_text"):
        parts += ["", "Clarifications Provided by User:", state["answers_text"]]
    if state.get("assumption_review"):
        parts += ["", "Assumption Decisions:", state["assumption_review"]]
    parts += ["", "Approved Campaign Briefing Document (where it differs from the brief above, the approved document wins):",
              bf.as_text(state.get("briefing") or {})]
    if feedback:
        parts += ["", "Refinement Requested by User:", feedback]
    return "\n".join(parts)


class _Agent:
    """One agent's record in the pipeline item, with its streamed reasoning."""

    def __init__(self, run: _Run, pipe: dict, name: str, message: str):
        self.run, self.pipe, self.name = run, pipe, name
        rec = next((a for a in pipe["agents"] if a["name"] == name), None)
        if rec is None:
            rec = {"name": name, "status": "active", "message": message, "result": None, "source": None,
                   "tokens": 0, "reasoning": "", "started_at": _now()}
            pipe["agents"].append(rec)
        else:
            rec.update(status="active", message=message)
        self.rec = rec
        self._buf = ""
        self._last = time.time()
        run.emit({"type": "agent_start", "item_id": pipe["id"], "agent": self._public()})

    def _public(self) -> dict:
        return {k: v for k, v in self.rec.items() if k != "reasoning"}

    def chunk(self, text: str) -> None:
        if len(self.rec["reasoning"]) < _REASONING_CAP:
            self.rec["reasoning"] += text
        self._buf += text
        if len(self._buf) >= 400 or time.time() - self._last >= 0.3:
            self.flush()

    def flush(self) -> None:
        if self._buf:
            self.run.emit({"type": "thought", "item_id": self.pipe["id"], "agent": self.name, "text": self._buf})
            self._buf = ""
        self._last = time.time()

    def mark(self, text: str) -> None:
        """A note in the agent's reasoning stream (e.g. why it fell back)."""
        self.chunk(f"\n\n[{text}]\n")

    def end(self, result: str, source: str, usage: dict | None = None, status: str = "complete") -> None:
        self.flush()
        self.rec.update(status=status, result=result, source=source, finished_at=_now(),
                        tokens=int(self.rec.get("tokens") or 0) + self.run.tokens(usage))
        self.run.emit({"type": "agent_done", "item_id": self.pipe["id"], "agent": self._public()})


def _ai_run(run: _Run, ag: _Agent, prompt: str, retries: int = 1) -> dict | None:
    """A streamed agent call with one reconnect retry; None when the model can't be used."""
    if not llm.available():
        return None
    for attempt in range(retries + 1):
        try:
            return agents.run(prompt, ag.chunk, run.stop)
        except llm.Cancelled:
            raise
        except Exception as exc:  # noqa: BLE001 -- retried, then the agent falls back
            print(f"[campaign-creator] {ag.name} attempt {attempt + 1} failed: {exc}")
            ag.mark(f"Model call failed: {exc}")
            if attempt < retries:
                time.sleep(2)
    return None


def build_blueprint(session_id: str, feedback: str | None, sf_conn: dict | None) -> dict:
    """Camille /flows/generate-stream on the approved brief; with `feedback`, a refinement of
    the current blueprint (the Architect edits the existing spec)."""
    sess = _load(session_id)
    st = sess["state"]
    if not st.get("briefing"):
        raise ValueError("Generate the Campaign Briefing Document first.")
    feedback = (feedback or "").strip() or None
    refine = bool(feedback and st.get("spec"))

    def body(run: _Run) -> None:
        state = run.state
        run.item({"kind": "user", "text": feedback or "Briefing Confirmed. Proceeding to Journey Blueprint Generation..."})
        pipe = run.item({"kind": "pipeline", "status": "running", "agents": [], "steps": [], "feedback": feedback,
                         "engine": llm.engine_label(), "refine": refine})
        prompt = compile_prompt(state, feedback)
        run.progress(pipe, "init", "System", "Initializing Agent Protocol...")
        metadata, meta_source = None, "standard"
        if sf_conn:
            try:
                metadata = salesforce.org_metadata(sf_conn, prompt)
                meta_source = "org"
                run.progress(pipe, "metadata", "Salesforce", f"Read {len(metadata)} objects from the connected org.")
            except Exception as exc:  # noqa: BLE001 -- the standard schema stands in
                run.progress(pipe, "metadata", "Salesforce", f"Couldn't read the org's objects ({exc}); using the standard schema.")
        if metadata is None:
            metadata = salesforce.standard_metadata()
            if not sf_conn:
                run.progress(pipe, "metadata", "Salesforce", "No org connected: mapping onto the standard Salesforce objects.")
        sources: dict[str, str] = {}
        usage_by: dict[str, dict] = {}
        briefing = state["briefing"]

        # 1. Document Analyst
        ag = _Agent(run, pipe, agents.ANALYST, "Reading campaign brief and extracting business requirements...")
        res = _ai_run(run, ag, agents.analyst_prompt(prompt))
        analysis = res["parsed"] if res and isinstance(res.get("parsed"), dict) else None
        if analysis is None:
            if res:
                ag.mark("Couldn't read the analysis; using the approved briefing's journey instead")
            analysis = agents.analysis_from_briefing(briefing)
        sources[agents.ANALYST] = "ai" if res and isinstance(res.get("parsed"), dict) else "rules"
        usage_by["Analyst"] = (res or {}).get("usage") or {}
        ag.end(analysis.get("businessGoal") or "Requirements extracted", sources[agents.ANALYST], usage_by["Analyst"])

        # 2. Salesforce Architect
        ag = _Agent(run, pipe, agents.ARCHITECT, "Mapping business requirements to Salesforce Flow features...")
        current = state.get("spec") if refine else None
        res = _ai_run(run, ag, agents.architect_prompt(analysis, metadata, current, feedback if refine else None))
        architect = agents.normalize_spec_envelope(res["parsed"], "rationale", "Mapped requirements to Salesforce Flow") if res else None
        sources[agents.ARCHITECT] = "ai" if architect is not None else "rules"
        if architect is None:
            if res:
                ag.mark("Couldn't read the flow specification; building it from the approved briefing")
            architect = {"rationale": "Flow built from the approved Campaign Briefing Document's journey flow.",
                         "spec": current or agents.spec_from_briefing(briefing)}
        usage_by["Architect"] = (res or {}).get("usage") or {}
        ag.end("Flow specification drafted", sources[agents.ARCHITECT], usage_by["Architect"])

        # 2.5 Flow QA Tester
        ag = _Agent(run, pipe, agents.QA, "Auditing and certifying Salesforce Flow specification for production readiness...")
        raw_spec = architect.get("spec")
        res = _ai_run(run, ag, agents.qa_prompt(analysis, raw_spec, metadata))
        qa = agents.normalize_spec_envelope(res["parsed"], "testReport", {"status": "PASSED_AND_CERTIFIED"}) if res else None
        certified = (qa or {}).get("spec") or raw_spec
        qa_report = (qa or {}).get("testReport") if isinstance((qa or {}).get("testReport"), dict) else None
        sources[agents.QA] = "ai" if qa else "rules"
        usage_by["FlowQATester"] = (res or {}).get("usage") or {}
        if qa_report:
            qa_summary = f"{qa_report.get('status') or 'Certified'}" + (f" · {qa_report.get('journeyCoveragePercent')}% coverage"
                                                                       if qa_report.get("journeyCoveragePercent") is not None else "")
        else:
            qa_summary = "Spec validated (no QA report)"
        ag.end(qa_summary, sources[agents.QA], usage_by["FlowQATester"])
        flow_xml.validate_flow_spec(certified)

        # 3. Visual Designer <-> Tester Agent (self-correction, up to 2 fixes)
        code, m_feedback, attempt, valid = "", "", 0, False
        designer_usage: dict = {}
        tester_usage: dict = {}
        redraw = False  # the last reply was cut off or unreadable, not a syntax error
        while attempt <= 2 and not valid:
            ag = _Agent(run, pipe, agents.DESIGNER, "Drawing Mermaid architecture diagram..." if attempt == 0
                        else f"Redrawing a more compact diagram (Attempt {attempt})..." if redraw
                        else f"Fixing Mermaid syntax errors (Attempt {attempt})...")
            res = _ai_run(run, ag, agents.designer_prompt(certified, m_feedback, prompt))
            drawn = None
            if res and isinstance(res.get("parsed"), dict) and res["parsed"].get("mermaid"):
                drawn = mermaid.sanitize(str(res["parsed"]["mermaid"]))
            designer_usage = res.get("usage") if res else designer_usage
            if not drawn and res and attempt < 2:
                # Not in Camille (it stopped with "Failed to generate visual diagram"): one more
                # try, asked to fit, before the rules diagram.
                cut = llm.truncated(res.get("usage"))
                ag.mark("The diagram was cut off at the output limit; asking for a more compact one" if cut
                        else "Couldn't read the diagram; asking for it again")
                ag.end("Diagram cut off" if cut else "Diagram unreadable", "ai", designer_usage, status="error")
                m_feedback = ("Your previous answer was cut off at the output limit before the JSON finished. Draw the "
                              "COMPLETE journey again but more compactly: short labels (at most 2 lines each), fold "
                              "repeated identical steps into one node per stage, and keep the JSON on as few lines as "
                              "possible. The JSON must be complete and valid.") if cut else (
                              "Your previous answer did not contain a readable JSON object with a \"mermaid\" key after "
                              "the ---JSON_START--- line. Output the full diagram again in exactly that format.")
                redraw = True
                attempt += 1
                continue
            if not drawn:
                if res:
                    ag.mark("Couldn't read the diagram; drawing it from the flow specification")
                code = agents.diagram_fallback(certified)
                sources[agents.DESIGNER] = "rules"
                ag.end("Diagram drawn from the flow specification", "rules", designer_usage)
                run.emit({"type": "mermaid", "item_id": pipe["id"], "mermaid": code})
                tag = _Agent(run, pipe, agents.TESTER, "Validating Mermaid syntax for render-ability...")
                tag.end("Skipped (diagram drawn by rules)", "rules", status="skipped")
                sources[agents.TESTER] = "rules"
                break
            code = drawn
            sources[agents.DESIGNER] = "ai"
            ag.end("Diagram Drafted" if attempt == 0 else "Diagram Fixed", "ai", designer_usage)
            run.emit({"type": "mermaid", "item_id": pipe["id"], "mermaid": code})
            tag = _Agent(run, pipe, agents.TESTER, "Validating Mermaid syntax for render-ability...")
            tres = _ai_run(run, tag, agents.tester_prompt(code), retries=0)
            if not tres:
                tag.end("Tester Error", "rules", status="error")
                sources[agents.TESTER] = "rules"
                break
            report = tres["parsed"] if isinstance(tres.get("parsed"), dict) else {"isValid": True, "feedback": ""}
            tester_usage = tres.get("usage") or {}
            raw_valid = report.get("isValid", True)
            valid = raw_valid if isinstance(raw_valid, bool) else str(raw_valid).strip().lower() != "false"
            sources[agents.TESTER] = "ai"
            if valid:
                tag.end("Syntax Validated ✅", "ai", tester_usage)
            else:
                tag.end("Syntax Error Detected ❌", "ai", tester_usage)
                m_feedback = str(report.get("feedback") or "Syntax invalid, please fix the diagram according to Mermaid rules.")
                redraw = False
                attempt += 1
                if attempt <= 2:
                    run.progress(pipe, "designer", "Self-Correction", "Routing feedback back to Designer...")
        usage_by["Designer"], usage_by["Tester"] = designer_usage, tester_usage

        # 4. Flow Validator
        ag = _Agent(run, pipe, agents.VALIDATOR, "Validating architecture against Salesforce limits...")
        res = _ai_run(run, ag, agents.validator_prompt(certified))
        if res:
            validation = res["parsed"] if isinstance(res.get("parsed"), dict) else {
                "isValid": True, "warnings": ["Validation parsing failed"], "recommendations": []}
            sources[agents.VALIDATOR] = "ai"
        else:
            validation = agents.validate_rules(certified, prompt)
            sources[agents.VALIDATOR] = "rules"
        usage_by["Validator"] = (res or {}).get("usage") or {}
        ag.end("Passed" if validation.get("isValid") not in (False, "false") else "Warnings Found",
               sources[agents.VALIDATOR], usage_by["Validator"])

        # 5. Technical Writer
        ag = _Agent(run, pipe, agents.WRITER, "Authoring final Technical Design documentation...")
        res = _ai_run(run, ag, agents.writer_prompt(analysis, certified, validation))
        if res and isinstance(res.get("parsed"), dict):
            docs = res["parsed"]
            sources[agents.WRITER] = "ai"
            writer_note = "Documentation Complete"
        else:
            docs = agents.docs_from(briefing, certified, validation)
            sources[agents.WRITER] = "rules"
            writer_note = "Documentation drafted from the brief"
        usage_by["Writer"] = (res or {}).get("usage") or {}
        ag.end(writer_note, sources[agents.WRITER], usage_by["Writer"])

        total = sum(int((u or {}).get("totalTokenCount") or 0) for u in usage_by.values())
        blueprint = {
            "rationale": architect.get("rationale") or analysis.get("businessGoal"),
            "planningSummary": analysis,
            "technicalDesign": {**architect, "spec": certified},
            "flowQaReport": qa_report or {"status": "NOT_RUN"},
            "validation": validation,
            "documentation": docs,
            "spec": certified,
            "mermaid": code,
            "tokenUsage": {"totalTokens": total, "breakdown": usage_by},
            "sources": sources,
            "metadataSource": meta_source,
            "metadataObjects": sorted(metadata.keys()),
            "specSummary": flow_xml.summarize(certified),
            "generatedAt": _now(),
            "briefingVersion": state.get("briefing_version"),
            "feedback": feedback,
            "engine": llm.engine_label(),
        }
        pipe["status"] = "done"
        # Status only: the browser already holds every agent's streamed reasoning.
        run.emit({"type": "item_status", "item_id": pipe["id"], "status": "done"})
        state.update(blueprint=blueprint, spec=certified, stage="blueprint", blueprint_stale=False,
                     blueprint_briefing_version=state.get("briefing_version"), deployment=None)
        run.item({"kind": "agent", "text": "The journey blueprint is ready for review: solution plan, technical design, "
                                            "QA report, validation and the journey flow diagram. Approve & Deploy when it looks right, "
                                            "or describe a change to refine it."})
        if run.commit():
            run.emit({"type": "blueprint", "blueprint": blueprint})
            run.done()

    return _start(session_id, "blueprint", body)


# ------------------------------------------------------------------ deploy --------------

def flow_label(state: dict) -> str:
    name = ((state.get("briefing") or {}).get("campaignName") or "").strip()
    return (name or "AI Generated Flow")[:80]


def flow_package(session_id: str) -> tuple[str, bytes]:
    """The deployable zip for the current spec -> (file name, bytes)."""
    state = _load(session_id)["state"]
    if not state.get("spec"):
        raise ValueError("Generate the journey blueprint first.")
    dep = state.get("deployment") or {}
    name = dep.get("flow_name") or flow_xml.flow_api_name(int(time.time() * 1000))
    xml = flow_xml.generate_flow_xml(state["spec"], name, flow_label(state))
    return f"{name}.zip", flow_xml.build_package(name, xml)


def deploy(session_id: str, sf_conn: dict | None) -> dict:
    """Camille /flows/deploy with its progress timeline: build the XML, push it, wait."""
    sess = _load(session_id)
    if not sess["state"].get("spec"):
        raise ValueError("Generate the journey blueprint first.")
    if not sf_conn:
        raise ValueError("Connect a Salesforce org before deploying.")

    def body(run: _Run) -> None:
        state = run.state
        name = flow_xml.flow_api_name(int(time.time() * 1000))
        dep = {"flow_name": name, "label": flow_label(state), "status": "running", "steps": [], "flow_url": None,
               "deploy_id": None, "errors": [], "started_at": _now(), "instance_url": sf_conn.get("instance_url")}

        def step(title: str, status: str, details: str) -> None:
            dep["steps"].append({"title": title, "status": status, "details": details})
            run.emit({"type": "deploy", "deployment": copy.deepcopy(dep)})

        def settle(status: str, details: str | None = None) -> None:
            if dep["steps"]:
                dep["steps"][-1]["status"] = status
                if details:
                    dep["steps"][-1]["details"] = details
            run.emit({"type": "deploy", "deployment": copy.deepcopy(dep)})

        step("Building Flow XML", "active", "Converting JSON specification to Salesforce Flow XML format...")
        try:
            xml = flow_xml.generate_flow_xml(state["spec"], name, dep["label"])
            package = flow_xml.build_package(name, xml)
            settle("complete", "Flow XML generated successfully")
            step("Pushing to Salesforce", "active", "Deploying flow via Metadata API...")
            waiting = {"shown": False}

            def on_status(sf_state: str, deploy_id: str) -> None:
                dep["deploy_id"] = deploy_id
                if not waiting["shown"]:
                    waiting["shown"] = True
                    settle("complete", "Deployment accepted by Salesforce")
                    step("Waiting for Salesforce", "active", f"Salesforce is processing the deployment... ({sf_state})")
                else:
                    dep["steps"][-1]["details"] = f"Salesforce is processing the deployment... ({sf_state})"
                    run.emit({"type": "deploy", "deployment": copy.deepcopy(dep)})

            result = salesforce.deploy(sf_conn, package, on_status=on_status)
            if result["success"]:
                settle("complete")
                dep["flow_url"] = salesforce.flows_url(sf_conn)
                dep["status"] = "succeeded"
                step("Flow Deployed Successfully", "complete", f"Flow Name: {name}")
                text = f"Deployed **{name}** to Salesforce as a Draft flow."
            else:
                dep["errors"] = result["errors"]
                dep["status"] = "failed"
                settle("error")
                step("Deployment Failed", "error", "; ".join(result["errors"]) or f"Salesforce reported {result['status']}")
                text = "The deployment failed. The Deployment panel lists what Salesforce reported."
        except (salesforce.SalesforceError, flow_xml.FlowError, OSError, ValueError) as exc:
            dep["status"] = "failed"
            dep["errors"] = [str(exc)]
            settle("error")
            step("Deployment Failed", "error", str(exc))
            text = f"The deployment failed: {exc}"
        except Exception as exc:  # noqa: BLE001 -- e.g. network errors from requests
            dep["status"] = "failed"
            dep["errors"] = [str(exc)]
            settle("error")
            step("Deployment Failed", "error", str(exc))
            text = f"The deployment failed: {exc}"
        dep["finished_at"] = _now()
        state["deployment"] = dep
        run.item({"kind": "agent", "text": text})
        if run.commit():
            run.done()

    return _start(session_id, "deploy", body)


# ------------------------------------------------------------------ reset / cancel ------

def reset(session_id: str) -> dict:
    """Start over: cancel any running step and clear the session (and its briefing versions).
    The linked artifact keeps its history; the next briefing becomes its next version."""
    sess = _load(session_id)
    jobs.cancel_active(session_id)
    with _lock(session_id):
        store.save_state(session_id, new_state(bool(sess["state"].get("auto_assume"))))
        store.clear_briefing_versions(session_id)
    return get_view(session_id)


def cancel(session_id: str) -> dict:
    _load(session_id)
    jobs.cancel_active(session_id)
    return get_view(session_id)


# ------------------------------------------------------------------ v3 artifact sync ----

def _spine_values(b: dict) -> dict:
    """The approved briefing's answers to the campaign decision spine (config/frameworks/
    campaign_spine.json), so the Flow Planner and Briefing Agent start from this plan."""
    b = bf.normalize_briefing(b)
    ov = b["campaignOverview"]
    segs = b["audienceSegmentation"]["segments"]
    vals = {
        "S1.2": ov.get("objective"),
        "S2.2": ov.get("primaryAudience") or (segs[0]["name"] if segs else ""),
        "S2.3": "; ".join(b["journeyEntryCriteria"]),
        "S4.1": ov.get("complianceRole"),
        "S5.2": segs[0]["messageFocus"] if segs else "",
        "S6.2": ov.get("primaryChannel"),
        "S8.0": "; ".join(f"{m['tier']}: {m['metrics']}" for m in b["measurementReporting"]),
    }
    return {k: v for k, v in vals.items() if v and v != bf.NOT_SPECIFIED}


def _sync_artifact(session_id: str) -> None:
    """Publish the briefing as this campaign's "campaign-planner" artifact (a new version per
    briefing version): it lists under My work and seeds the downstream agents' workspaces."""
    try:
        sess = _load(session_id)
        b = sess["state"].get("briefing")
        if not b or not sess["brand"]:
            return
        from strategy import agent_forms, hierarchy, v3_artifacts
        plan = campaign = None
        try:
            tree = hierarchy.tree(sess["brand"])
            plan = next((p for p in tree["engagement_plans"] if p["id"] == sess["plan_id"]), None)
            campaign = next((c for c in (plan or {}).get("campaigns", []) if c["id"] == sess["campaign_id"]), None)
        except Exception:  # noqa: BLE001 -- an unknown brand still gets an artifact
            pass
        form = agent_forms.build_cards("campaign-planner", sess["brand"], plan, campaign)
        inputs = {}
        for st in form.get("stages", []):
            for p in st["data_points"]:
                inputs[p["key"]] = {"label": p["label"], "derivation": p["derivation"],
                                    "value": p["value"] if p["derivation"] == "derive" else None,
                                    "confirmed": p["derivation"] != "confirm"}
        for key, value in _spine_values(b).items():
            if key in inputs:
                inputs[key]["value"] = value
                inputs[key]["confirmed"] = True
        version = sess["state"].get("briefing_version")
        extras = [{"label": "Campaign Briefing Document",
                   "value": f"{b.get('campaignName') or 'Campaign'} — briefing v{version}, built in the Campaign Planner."}]
        art = v3_artifacts.generate("campaign-planner", sess["brand"], sess["plan_id"], sess["campaign_id"],
                                    sess["title"] or "Campaign plan", inputs, extras)
        if art and art.get("id") != sess["artifact_id"]:
            store.set_fields(session_id, artifact_id=art["id"])
    except Exception as exc:  # noqa: BLE001 -- publishing is a side effect; never break the planner over it
        print(f"[campaign-creator] artifact sync skipped: {exc}")
