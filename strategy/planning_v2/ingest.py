"""Stage 1 — Ingest.

Parses a PDF, PPTX, or DOCX strategic-plan document into `IngestedBlock`s that
each carry a page or section reference, so every fact stage 2 (Extract) pulls
out can be traced back to where it came from (the brief's provenance
requirement starts here, not at extraction).

This is a different, purpose-built parser from `strategy/document_intake.py`
(the existing chat-upload path): that one flattens everything to a single
~8000-char string with no structure and no PPTX support, which is fine for a
quick chat attachment but cannot carry per-fact provenance or handle a full
strategic-plan deck. It stays untouched; this module is new stage-1 code for
the planning engine specifically.
"""
from __future__ import annotations

import io
from pathlib import Path
from typing import Literal

from pydantic import BaseModel, Field

DocFormat = Literal["pdf", "pptx", "docx"]
_SUPPORTED = {"pdf", "pptx", "docx"}


class IngestedBlock(BaseModel):
    """One extracted unit of text plus where it came from."""

    source_id: str  # e.g. "p.3", "slide.5", "section.Market Landscape"
    order: int
    heading: str = ""
    text: str


class IngestedDocument(BaseModel):
    filename: str
    format: DocFormat
    blocks: list[IngestedBlock] = Field(default_factory=list)

    def tagged_text(self, max_chars: int | None = None) -> str:
        """Blocks joined with an inline `[source_id]` tag on each, which is
        what stage 2's extraction prompt shows the LLM so it can cite a
        source_id per fact it pulls out."""
        parts = [f"[{b.source_id}]\n{b.text}" for b in self.blocks if b.text.strip()]
        text = "\n\n".join(parts)
        if max_chars and len(text) > max_chars:
            text = text[:max_chars] + "\n\n[...truncated...]"
        return text


def _detect_format(filename: str) -> DocFormat:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in _SUPPORTED:
        raise ValueError(f"Unsupported file type: .{ext or 'unknown'} (supported: .pdf, .pptx, .docx)")
    return ext  # type: ignore[return-value]


def ingest_bytes(content: bytes, filename: str) -> IngestedDocument:
    fmt = _detect_format(filename)
    if fmt == "pdf":
        blocks = _ingest_pdf(content)
    elif fmt == "pptx":
        blocks = _ingest_pptx(content)
    else:
        blocks = _ingest_docx(content)
    return IngestedDocument(filename=filename, format=fmt, blocks=blocks)


def ingest_path(path: str | Path) -> IngestedDocument:
    p = Path(path)
    return ingest_bytes(p.read_bytes(), p.name)


def _ingest_pdf(content: bytes) -> list[IngestedBlock]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    blocks: list[IngestedBlock] = []
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        if not text:
            continue
        blocks.append(IngestedBlock(source_id=f"p.{i}", order=i, text=text))
    return blocks


def _ingest_pptx(content: bytes) -> list[IngestedBlock]:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(content))
    blocks: list[IngestedBlock] = []
    for i, slide in enumerate(prs.slides, start=1):
        chunks: list[str] = []
        title = ""
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            frame_text = "\n".join(
                para.text for para in shape.text_frame.paragraphs if para.text.strip()
            ).strip()
            if not frame_text:
                continue
            if shape == slide.shapes.title:
                title = frame_text
            chunks.append(frame_text)
        # also pick up notes, since strategic rationale often lives in speaker notes
        if slide.has_notes_slide:
            notes = (slide.notes_slide.notes_text_frame.text or "").strip()
            if notes:
                chunks.append(f"[speaker notes] {notes}")
        text = "\n".join(chunks).strip()
        if not text:
            continue
        blocks.append(IngestedBlock(source_id=f"slide.{i}", order=i, heading=title, text=text))
    return blocks


def _ingest_docx(content: bytes) -> list[IngestedBlock]:
    from docx import Document

    doc = Document(io.BytesIO(content))
    blocks: list[IngestedBlock] = []
    heading = "Untitled"
    order = 0
    buffer: list[str] = []

    def _flush() -> None:
        nonlocal buffer, order
        text = "\n".join(buffer).strip()
        if text:
            order += 1
            blocks.append(IngestedBlock(source_id=f"section.{heading}", order=order, heading=heading, text=text))
        buffer = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "") if para.style else ""
        if style.startswith("Heading") or style == "Title":
            _flush()
            heading = text
            continue
        buffer.append(text)
    _flush()
    return blocks
